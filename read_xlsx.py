import openpyxl, sys

def dump(path):
    print("=" * 70)
    print("FILE:", path)
    print("=" * 70)
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        print(f"\n### SHEET: {ws.title}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
        rows = list(ws.iter_rows(values_only=True))
        for i, r in enumerate(rows[:20], 1):
            print(f"  R{i}:", r)

for p in sys.argv[1:]:
    dump(p)
