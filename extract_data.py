import openpyxl, json, datetime

def clean(v):
    if v is None: return None
    if isinstance(v, datetime.datetime): return v.strftime("%Y-%m-%d")
    return v

# --- DBWB 9月 sheet: 设备名称, 耗材名称, 库存数量, 备注 ---
wb1 = openpyxl.load_workbook("D:/RjDir/UserData/Desktop/DBWB耗材盘点.xlsx", data_only=True)
ws1 = wb1["9月"]
dbwb = []
current_equip = None
for row in ws1.iter_rows(min_row=2, values_only=True):
    equip, name, qty, note = row[0], row[1], row[2], row[3]
    if equip: current_equip = clean(equip)
    if name is None: continue
    dbwb.append({
        "process": "焊线工序(DBWB)",
        "equipment": current_equip,
        "name": clean(name),
        "stock": int(qty) if qty is not None else 0,
        "note": clean(note) or ""
    })

# --- 耦合段配件统计: 序号, 配件名称, 数量, 备注 ---
wb2 = openpyxl.load_workbook("D:/RjDir/UserData/Downloads/耦合段配件统计.xlsx", data_only=True)
ws2 = wb2["配件统计"]
coupling = []
for row in ws2.iter_rows(min_row=3, values_only=True):
    seq, name, qty, note = row[0], row[1], row[2], row[3]
    if name is None: continue
    coupling.append({
        "process": "耦合段",
        "equipment": "耦合段机台",
        "name": clean(name),
        "stock": int(qty) if qty is not None else 0,
        "note": clean(note) or ""
    })

print("DBWB rows:", len(dbwb))
for d in dbwb[:3]: print("  ", d)
print("Coupling rows:", len(coupling))
for c in coupling[:3]: print("  ", c)

out = {"dbwb": dbwb, "coupling": coupling}
with open("C:/Users/admin/WorkBuddy/2026-09-17-11-26-50/seed_data.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=2)
print("saved seed_data.json")
